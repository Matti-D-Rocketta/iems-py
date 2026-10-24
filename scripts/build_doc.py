"""Build docs/iEMS_py_documentation.docx (and .pdf) from docs/src/iems_documentation.md.

Usage: python scripts/build_doc.py [--no-pdf]

Supported Markdown subset: ``#``/``##``/``###`` headings (auto-numbered), paragraphs, ``- ``
bullets, ``1. `` numbered items, pipe tables (first row = header), ``![caption](path)`` images
(auto-numbered captions), ``> `` call-out boxes, fenced code blocks, ``---`` page breaks, and
inline ``**bold**``, ``*italic*``, `` `code` ``. The PDF is exported through Microsoft Word
(COM automation) when available, with the table of contents refreshed.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "src" / "iems_documentation.md"
OUT_DOCX = ROOT / "docs" / "iEMS_py_documentation.docx"
OUT_PDF = ROOT / "docs" / "iEMS_py_documentation.pdf"

FONT = "Calibri"
MONO = "Consolas"
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
SHADE = "EEF3FB"
SHADE_CODE = "F3F2EF"
HEADER_SHADE = "DCE8F8"

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


# --------------------------------------------------------------------------------------
# Low-level helpers
# --------------------------------------------------------------------------------------
def shade(cell_or_par: object, hex_fill: str) -> None:
    """Background colour for a table cell or a paragraph."""
    el = (
        cell_or_par._tc.get_or_add_tcPr()
        if hasattr(cell_or_par, "_tc")
        else cell_or_par._p.get_or_add_pPr()
    )
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    el.append(shd)


def add_field(paragraph: object, instr: str) -> None:
    """Insert a Word field (TOC, PAGE…) into ``paragraph``."""
    run = paragraph.add_run()
    for tag, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        run._r.append(el)


def add_inline(
    paragraph: object, text: str, size: float | None = None, color: RGBColor | None = None
) -> None:
    """Add ``text`` with **bold**, *italic* and `code` spans."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
            run.font.size = Pt((size or 10.5) - 1)
            run.font.color.rgb = RGBColor(0x1C, 0x5C, 0xAB)
        elif part.startswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size and not part.startswith("`"):
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def set_cell_borders(table: object) -> None:
    """Light horizontal rules only."""
    tbl = table._tbl
    pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "C9C7C1")
        borders.append(b)
    for edge in ("left", "right", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "nil")
        borders.append(b)
    pr.append(borders)


# --------------------------------------------------------------------------------------
# Document builder
# --------------------------------------------------------------------------------------
class Builder:
    def __init__(self) -> None:
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.num = [0, 0, 0]
        self._styles()
        self._page()

    # --- setup -----------------------------------------------------------------------
    def _styles(self) -> None:
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, before in (
            ("Heading 1", 18, 18),
            ("Heading 2", 14, 14),
            ("Heading 3", 11.5, 10),
        ):
            h = st[name]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = ACCENT if name != "Heading 3" else INK
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            st[name].font.name = FONT
            st[name].font.size = Pt(10.5)
            st[name].paragraph_format.space_after = Pt(3)

    def _page(self) -> None:
        sec = self.doc.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        hp = sec.header.paragraphs[0]
        hp.text = "iEMS-py — Architecture and operation"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = INK2
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, "PAGE")
        fp.runs[0].font.size = Pt(8.5)

    # --- blocks ----------------------------------------------------------------------
    def title_page(self, title: str, subtitle: str, lines: list[str]) -> None:
        for _ in range(6):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        r = p.add_run(title)
        r.font.size = Pt(28)
        r.bold = True
        r.font.color.rgb = ACCENT
        p = self.doc.add_paragraph()
        r = p.add_run(subtitle)
        r.font.size = Pt(15)
        r.font.color.rgb = INK2
        self.doc.add_paragraph()
        for line in lines:
            p = self.doc.add_paragraph()
            add_inline(p, line, size=10.5, color=INK2)
        self.page_break()
        p = self.doc.add_paragraph()
        r = p.add_run("Contents")
        r.font.size = Pt(16)
        r.bold = True
        r.font.color.rgb = ACCENT
        add_field(self.doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
        self.page_break()

    def heading(self, level: int, text: str) -> None:
        if level == 1:
            self.num = [self.num[0] + 1, 0, 0]
            label = f"{self.num[0]}"
        elif level == 2:
            self.num[1] += 1
            self.num[2] = 0
            label = f"{self.num[0]}.{self.num[1]}"
        else:
            self.num[2] += 1
            label = f"{self.num[0]}.{self.num[1]}.{self.num[2]}"
        if level == 1 and self.num[0] > 1:
            self.page_break()
        self.doc.add_heading(f"{label}  {text.replace(chr(96), '')}", level=level)

    def paragraph(self, text: str) -> None:
        add_inline(self.doc.add_paragraph(), text)

    def bullet(self, text: str, numbered: bool = False) -> None:
        p = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        add_inline(p, text)

    def callout(self, lines: list[str]) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        shade(cell, SHADE)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        first = True
        for line in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            add_inline(p, line, size=10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def code(self, lines: list[str]) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        shade(cell, SHADE_CODE)
        first = True
        for line in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line if line else " ")
            r.font.name = MONO
            r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
            r.font.size = Pt(8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def table(self, rows: list[list[str]]) -> None:
        self.tab_n += 1
        ncol = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncol)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_borders(table)
        for i, row in enumerate(rows):
            for j in range(ncol):
                cell = table.cell(i, j)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                text = row[j] if j < len(row) else ""
                add_inline(p, text, size=9)
                if i == 0:
                    shade(cell, HEADER_SHADE)
                    for r in p.runs:
                        r.bold = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def image(self, path: Path, caption: str) -> None:
        self.fig_n += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(16.4))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(c, f"**Figure {self.fig_n}** — {caption}", size=9, color=INK2)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------------------
# Markdown subset parser
# --------------------------------------------------------------------------------------
def parse(md: str, b: Builder) -> None:
    lines = md.splitlines()
    i = 0
    para: list[str] = []

    def flush() -> None:
        if para:
            b.paragraph(" ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            b.code(block)
        elif stripped == "---":
            flush()
            b.page_break()
        elif stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            b.heading(min(level, 3), stripped[level:].strip())
        elif stripped.startswith("!["):
            flush()
            m = re.match(r"!\[(.*)\]\((.*)\)", stripped)
            assert m, stripped
            b.image(ROOT / m.group(2), m.group(1))
        elif stripped.startswith("> "):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            b.callout([s for s in block if s])
            continue
        elif stripped.startswith("|"):
            flush()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            b.table(rows)
            continue
        elif stripped.startswith("- "):
            flush()
            b.bullet(stripped[2:])
        elif re.match(r"\d+\. ", stripped):
            flush()
            b.bullet(re.sub(r"^\d+\. ", "", stripped), numbered=True)
        elif stripped == "":
            flush()
        else:
            para.append(line)
        i += 1
    flush()


def export_pdf(docx_path: Path, pdf_path: Path) -> bool:
    """Export through Word (updates the TOC and fields). Returns False if Word is unavailable."""
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError:
        return False
    try:
        word = win32com.client.DispatchEx("Word.Application")
    except Exception:
        return False
    try:
        word.Visible = False
        doc = word.Documents.Open(str(docx_path))
        doc.Fields.Update()
        for toc in doc.TablesOfContents:
            toc.Update()
        doc.Save()
        doc.ExportAsFixedFormat(str(pdf_path), 17)  # 17 = wdExportFormatPDF
        doc.Close(False)
        return True
    finally:
        word.Quit()


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args(argv)

    md = SRC.read_text(encoding="utf-8")
    head, _, body = md.partition("\n---\n")  # front matter: title / subtitle / lines
    meta = [s.strip() for s in head.strip().splitlines() if s.strip()]
    b = Builder()
    b.title_page(meta[0].lstrip("# ").strip(), meta[1], meta[2:])
    parse(body, b)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    b.doc.save(OUT_DOCX)
    print(f"docx : {OUT_DOCX}  ({b.fig_n} figures, {b.tab_n} tables)")
    if not args.no_pdf:
        if export_pdf(OUT_DOCX, OUT_PDF):
            print(f"pdf  : {OUT_PDF}")
        else:
            print("pdf  : Word not available here — open the .docx in Word and export to PDF")
    return 0


if __name__ == "__main__":
    sys.exit(main())
